"""Relatório diário de notas — lançadas (SISMOURA), a lançar e pendentes (manual ou por XML).

Uso:  python app.py   → abre http://localhost:3003 (porta e lojas visíveis definidas no .env)
O banco SISMOURA é acessado SOMENTE LEITURA. As notas a lançar/pendentes
ficam num SQLite local (dados/relatorio.db) e continuam em aberto até serem
lançadas no SISMOURA (detectado automaticamente) ou baixadas manualmente.
"""
import io
import json
import os
import re
import socket
import sqlite3
import sys
import threading
import webbrowser
import xml.etree.ElementTree as ET
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from flask import Flask, abort, jsonify, request, send_file, send_from_directory

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE / ".claude" / "skills" / "banco-sismoura" / "scripts"))
from consulta import LOJAS as TODAS_LOJAS, carregar_config, conectar  # noqa: E402

cfg = carregar_config()  # também carrega o .env
DB_LOCAL = BASE / "dados" / "relatorio.db"
PORTA = int(os.getenv("PORT", "3003"))


def lojas_permitidas():
    """LOJAS_PERMITIDAS do .env: 'todas' (todas as configuradas) ou lista, ex.: 'pontanegra'."""
    valor = os.getenv("LOJAS_PERMITIDAS", "todas").strip().lower()
    if valor in ("", "todas"):
        return {k: v for k, v in TODAS_LOJAS.items() if k in cfg["lojas"]}
    ids = [x.strip() for x in valor.split(",") if x.strip()]
    for x in ids:
        if x not in TODAS_LOJAS:
            sys.exit(f"ERRO: LOJAS_PERMITIDAS tem a loja '{x}', que não existe. Use: {', '.join(TODAS_LOJAS)}")
        if x not in cfg["lojas"]:
            sys.exit(f"ERRO: a loja '{x}' está em LOJAS_PERMITIDAS mas falta {x.upper()}_SERVIDOR no .env")
    return {x: TODAS_LOJAS[x] for x in TODAS_LOJAS if x in ids}


# Só as lojas deste servidor: todo o sistema (consultas, telas, notas manuais) usa este dicionário.
LOJAS = lojas_permitidas()
NOMES_PERMITIDOS = set(LOJAS.values())


def loja_ok(nome):
    """Nota sem loja definida é visível; com loja, só se for uma das permitidas."""
    return not nome or nome in NOMES_PERMITIDOS


def exigir_nota_permitida(con, id_):
    r = con.execute("SELECT loja FROM notas_manuais WHERE id = ?", (id_,)).fetchone()
    if r is None or not loja_ok(r["loja"]):
        abort(404)

# CNPJ do destinatário na NF-e → loja (conferido com os XMLs gravados em cada banco).
# Obs.: Flores também recebe notas no CNPJ 0001-30; nesses casos ajuste a loja na tela.
CNPJ_LOJA = {
    "19234190000130": "Parque Dez",
    "19234190000300": "Ponta Negra",
    "19234190000482": "Monte das Oliveiras",
    "19234190000563": "Flores",
    "19234190000644": "CD",
    "19234190000725": "Morada do Sol",
}

# Usuário que lançou: Operador da entrada → usuário que abriu → operador do título gerado no Contas a Pagar.
SQL_LANCADAS = """
SELECT e.Chave, e.Chave_NFE, e.Data_Cadastro, e.Nota_Fiscal, e.Serie, e.Nome_Fornecedor,
       e.Valor_Nota, e.Data_Emissao, u.Nome AS Usuario, b.Bonificado
FROM Entrada_Produto e
OUTER APPLY (SELECT SUM(i.Valor_Total) AS Bonificado
             FROM Entrada_Produto_Itens i
             WHERE i.Chave = e.Chave AND REPLACE(i.CFOP, '.', '') IN ('1910', '2910')) b
OUTER APPLY (SELECT TOP 1 NULLIF(cp.Operador, 0) AS Op
             FROM Contas_Pagar cp
             WHERE cp.Codigo_Entrada = e.Chave AND NULLIF(cp.Operador, 0) IS NOT NULL) cp
LEFT JOIN Usuario u ON u.Codigo = COALESCE(NULLIF(e.Operador, 0), NULLIF(e.Aberto_Usuario, 0), cp.Op)
WHERE e.Data_Cadastro >= ? AND e.Data_Cadastro < ?
ORDER BY e.Data_Cadastro
"""

# Contas a pagar geradas pelas notas lançadas no período (títulos excluídos ficam de fora).
SQL_TITULOS = """
SELECT cp.Codigo_Entrada, cp.Codigo, cp.Parcela, cp.Data_Vencimento, cp.Valor, cp.Pago
FROM Contas_Pagar cp
JOIN Entrada_Produto e ON e.Chave = cp.Codigo_Entrada
WHERE e.Data_Cadastro >= ? AND e.Data_Cadastro < ? AND cp.Data_Exclusao IS NULL
ORDER BY cp.Codigo_Entrada, cp.Data_Vencimento, cp.Parcela
"""

app = Flask(__name__, static_folder=None)
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024


# ---------- banco local (notas a lançar / pendentes) ----------

COLUNAS_EXTRAS = {
    "chave": "TEXT", "serie": "TEXT", "emissao": "TEXT", "cnpj_emit": "TEXT",
    "origem": "TEXT DEFAULT 'manual'", "resolvido_em": "TEXT",
    "lancada_em": "TEXT", "lancada_loja": "TEXT",
}


def local():
    DB_LOCAL.parent.mkdir(exist_ok=True)
    con = sqlite3.connect(DB_LOCAL)
    con.row_factory = sqlite3.Row
    con.execute("""CREATE TABLE IF NOT EXISTS notas_manuais (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        data TEXT NOT NULL,
        tipo TEXT NOT NULL CHECK (tipo IN ('alancar', 'pendente')),
        loja TEXT, nf TEXT, fornecedor TEXT, valor REAL DEFAULT 0, motivo TEXT,
        criado_em TEXT DEFAULT CURRENT_TIMESTAMP)""")
    existentes = {r[1] for r in con.execute("PRAGMA table_info(notas_manuais)")}
    for col, tipo in COLUNAS_EXTRAS.items():
        if col not in existentes:
            con.execute(f"ALTER TABLE notas_manuais ADD COLUMN {col} {tipo}")
    return con


# ---------- SISMOURA (somente leitura) ----------

def consultar_varias(loja, *consultas):
    """Roda várias consultas (sql, params) na mesma conexão. Devolve uma lista de linhas por consulta."""
    conn = conectar(cfg, loja, 60 if loja == "moradadosol" else 30)
    try:
        cur = conn.cursor()
        cur.execute("SET TRANSACTION ISOLATION LEVEL READ UNCOMMITTED")
        resultados = []
        for sql, params in consultas:
            cur.execute(sql, *params)
            cols = [c[0] for c in cur.description]
            resultados.append([dict(zip(cols, r)) for r in cur.fetchall()])
        return resultados
    finally:
        conn.rollback()
        conn.close()


def consultar(loja, sql, *params):
    return consultar_varias(loja, (sql, params))[0]


def em_paralelo(func, lojas):
    """Roda func(loja) em todas as lojas; devolve ({loja: resultado}, [erros])."""
    res, erros = {}, []
    with ThreadPoolExecutor(max_workers=max(1, len(lojas))) as ex:
        futuros = {l: ex.submit(func, l) for l in lojas}
        for l, f in futuros.items():
            try:
                res[l] = f.result()
            except Exception as e:  # loja fora do ar não derruba o relatório
                erros.append({"loja": LOJAS[l], "erro": str(e).split("]")[-1][:200]})
    return res, erros


def buscar_lancadas(loja, dia):
    ini = datetime.combine(dia, datetime.min.time())
    periodo = (ini, ini + timedelta(days=1))
    notas, titulos = consultar_varias(loja, (SQL_LANCADAS, periodo), (SQL_TITULOS, periodo))
    por_entrada = {}
    for t in titulos:
        por_entrada.setdefault(t["Codigo_Entrada"], []).append(t)
    out = []
    for d in notas:
        emissao = d["Data_Emissao"].date() if d["Data_Emissao"] else None
        out.append({
            "titulos": [{
                "codigo": t["Codigo"],
                "parcela": t["Parcela"],
                "vencimento": t["Data_Vencimento"].strftime("%d/%m/%Y"),
                "prazo": (t["Data_Vencimento"].date() - emissao).days if emissao else None,
                "valor": float(t["Valor"] or 0),
                "pago": t["Pago"] == "S",
            } for t in por_entrada.get(d["Chave"], [])],
            "loja": LOJAS[loja],
            "chave": d["Chave"],
            "chave_nfe": (d["Chave_NFE"] or "").strip(),
            "hora": d["Data_Cadastro"].strftime("%H:%M") if d["Data_Cadastro"] else "",
            "nf": (d["Nota_Fiscal"] or "").strip(),
            "serie": (d["Serie"] or "").strip(),
            "fornecedor": (d["Nome_Fornecedor"] or "").strip(),
            "valor": float(d["Valor_Nota"] or Decimal(0)),
            "bonificado": float(d["Bonificado"] or Decimal(0)),  # itens com CFOP 1.910/2.910 (bonificação)
            "emissao": d["Data_Emissao"].strftime("%d/%m/%Y") if d["Data_Emissao"] else "",
            "usuario": " ".join((d["Usuario"] or "Não identificado").split()).title(),
        })
    return out


def so_digitos(s):
    return re.sub(r"\D", "", str(s or ""))


def localizar_no_sismoura(chaves, nfs_por_loja):
    """Procura notas já lançadas em Entrada_Produto.
    chaves: chaves de acesso (procuradas em todas as lojas).
    nfs_por_loja: {nome_loja: {nf com 9 dígitos}} para notas digitadas sem chave.
    Retorna ({chave: (data, loja)}, {(loja, nf9): [datas]}, erros)."""
    chaves = sorted({c for c in chaves if len(c) == 44})
    nome_para_id = {v: k for k, v in LOJAS.items()}
    lojas = set(LOJAS) if chaves else {nome_para_id[n] for n in nfs_por_loja if n in nome_para_id}
    if not lojas:
        return {}, {}, []

    def busca(loja):
        conds = []
        if chaves:
            conds.append("Chave_NFE IN (%s)" % ",".join(f"'{c}'" for c in chaves))
        nfs = sorted(nfs_por_loja.get(LOJAS[loja], ()))
        if nfs:
            conds.append("Nota_Fiscal IN (%s)" % ",".join(f"'{n}'" for n in nfs))
        if not conds:
            return []
        return consultar(loja, "SELECT Chave_NFE, Nota_Fiscal, Data_Cadastro FROM Entrada_Produto WHERE "
                         + " OR ".join(conds))

    res, erros = em_paralelo(busca, sorted(lojas))
    por_chave, por_nf = {}, {}
    for loja, linhas in res.items():
        for r in linhas:
            dia = r["Data_Cadastro"].date().isoformat() if r["Data_Cadastro"] else None
            if r["Chave_NFE"] and r["Chave_NFE"] in chaves:
                por_chave.setdefault(r["Chave_NFE"], (dia, LOJAS[loja]))
            nf9 = so_digitos(r["Nota_Fiscal"]).zfill(9)
            por_nf.setdefault((LOJAS[loja], nf9), []).append(dia)
    return por_chave, por_nf, erros


def atualizar_lancadas(con, rows):
    """Marca lancada_em nas notas manuais que já aparecem no SISMOURA."""
    abertas = [r for r in rows if not r["lancada_em"] and not r["resolvido_em"]]
    if not abertas:
        return []
    chaves = [r["chave"] for r in abertas if r["chave"]]
    nfs = {}
    for r in abertas:
        if not r["chave"] and so_digitos(r["nf"]):
            nfs.setdefault(r["loja"], set()).add(so_digitos(r["nf"]).zfill(9)[-9:])
    por_chave, por_nf, erros = localizar_no_sismoura(chaves, nfs)
    for r in abertas:
        achou = None
        if r["chave"] and r["chave"] in por_chave:
            achou = por_chave[r["chave"]]
        elif not r["chave"] and so_digitos(r["nf"]):
            datas = [d for d in por_nf.get((r["loja"], so_digitos(r["nf"]).zfill(9)[-9:]), []) if d and d >= r["data"]]
            if datas:
                achou = (min(datas), r["loja"])
        if achou and achou[0]:
            con.execute("UPDATE notas_manuais SET lancada_em = ?, lancada_loja = ? WHERE id = ?",
                        (achou[0], achou[1], r["id"]))
    return erros


# ---------- leitura de XML de NF-e ----------

def _sem_ns(el):
    for e in el.iter():
        if isinstance(e.tag, str) and "}" in e.tag:
            e.tag = e.tag.split("}", 1)[1]
    return el


def _txt(el, caminho):
    achado = el.find(caminho)
    return (achado.text or "").strip() if achado is not None and achado.text else ""


def ler_nfe(conteudo):
    """Lê um XML de NF-e (nfeProc ou NFe). Retorna dict ou levanta ValueError."""
    try:
        raiz = _sem_ns(ET.fromstring(conteudo))
    except ET.ParseError:
        raise ValueError("XML inválido")
    inf = raiz.find(".//infNFe")
    if inf is None:
        raise ValueError("não é uma NF-e (pode ser CT-e ou evento)")
    chave = so_digitos(inf.get("Id")) or so_digitos(_txt(raiz, ".//protNFe/infProt/chNFe"))
    emit_doc = _txt(inf, "emit/CNPJ") or _txt(inf, "emit/CPF")
    return {
        "chave": chave,
        "nf": _txt(inf, "ide/nNF"),
        "serie": _txt(inf, "ide/serie"),
        "emissao": (_txt(inf, "ide/dhEmi") or _txt(inf, "ide/dEmi"))[:10],
        "fornecedor": _txt(inf, "emit/xNome"),
        "cnpj_emit": emit_doc,
        "cnpj_dest": _txt(inf, "dest/CNPJ") or _txt(inf, "dest/CPF"),
        "valor": float(_txt(inf, "total/ICMSTot/vNF") or 0),
    }


def extrair_xmls(nome, dados, profundidade=0):
    """Gera (nome, bytes) de cada .xml, abrindo .zip (inclusive zip dentro de zip)."""
    if nome.lower().endswith(".zip") and profundidade < 3:
        with zipfile.ZipFile(io.BytesIO(dados)) as zf:
            for n in zf.namelist():
                if not n.endswith("/"):
                    yield from extrair_xmls(n, zf.read(n), profundidade + 1)
    elif nome.lower().endswith(".xml"):
        yield nome, dados


# ---------- API ----------

@app.get("/api/lancadas")
def api_lancadas():
    dia = date.fromisoformat(request.args.get("data", date.today().isoformat()))
    lojas = [l for l in request.args.get("lojas", ",".join(LOJAS)).split(",") if l in LOJAS]
    res, erros = em_paralelo(lambda l: buscar_lancadas(l, dia), lojas)
    notas = sorted((n for v in res.values() for n in v), key=lambda n: n["hora"])
    return jsonify({"data": dia.isoformat(), "notas": notas, "erros": erros})


@app.get("/api/lojas")
def api_lojas():
    return jsonify([{"id": k, "nome": v} for k, v in LOJAS.items()])


SQL_MANUAIS_DIA = """SELECT * FROM notas_manuais
    WHERE data <= :d AND IFNULL(resolvido_em, '9999') >= :d AND IFNULL(lancada_em, '9999') >= :d
    ORDER BY data, id"""


@app.get("/api/manuais")
def api_manuais_listar():
    """Notas manuais em aberto na data (chegadas até a data e ainda não lançadas/baixadas),
    mais as que foram lançadas/baixadas exatamente na data (para exibir o ✓)."""
    d = request.args["data"]

    def do_dia(con):
        return [r for r in con.execute(SQL_MANUAIS_DIA, {"d": d}).fetchall() if loja_ok(r["loja"])]

    with local() as con:
        erros = atualizar_lancadas(con, do_dia(con))
        rows = do_dia(con)
    return jsonify({"notas": [dict(r) for r in rows], "erros": erros})


CAMPOS_EDITAVEIS = ("tipo", "loja", "nf", "serie", "fornecedor", "valor", "motivo", "emissao", "resolvido_em")


@app.post("/api/manuais")
def api_manuais_criar():
    b = request.get_json()
    if not loja_ok(b.get("loja")):
        abort(403)
    with local() as con:
        cur = con.execute(
            "INSERT INTO notas_manuais (data, tipo, loja, nf, fornecedor, valor, motivo, origem) VALUES (?,?,?,?,?,?,?, 'manual')",
            (b["data"], b["tipo"], b.get("loja"), b.get("nf"), b.get("fornecedor"),
             float(b.get("valor") or 0), b.get("motivo")))
        row = con.execute("SELECT * FROM notas_manuais WHERE id = ?", (cur.lastrowid,)).fetchone()
    return jsonify(dict(row))


@app.put("/api/manuais/<int:id_>")
def api_manuais_editar(id_):
    b = request.get_json()
    campos = [c for c in CAMPOS_EDITAVEIS if c in b]
    if "loja" in b and not loja_ok(b["loja"]):
        abort(403)
    if campos:
        vals = [float(b[c] or 0) if c == "valor" else (b[c] or None) for c in campos]
        with local() as con:
            exigir_nota_permitida(con, id_)
            con.execute(f"UPDATE notas_manuais SET {', '.join(c + ' = ?' for c in campos)} WHERE id = ?",
                        (*vals, id_))
    return jsonify({"ok": True})


@app.delete("/api/manuais/<int:id_>")
def api_manuais_excluir(id_):
    with local() as con:
        exigir_nota_permitida(con, id_)
        con.execute("DELETE FROM notas_manuais WHERE id = ?", (id_,))
    return jsonify({"ok": True})


@app.post("/api/importar")
def api_importar():
    """Importa XMLs de NF-e (um ou vários, soltos ou em .zip) como notas a lançar ou pendentes."""
    tipo = request.form.get("tipo", "alancar")
    dia = request.form["data"]
    loja_forcada = request.form.get("loja") or None
    motivo = request.form.get("motivo") or None
    if not loja_ok(loja_forcada):
        abort(403)

    notas, ignorados = {}, []
    for f in request.files.getlist("arquivos"):
        try:
            for nome, conteudo in extrair_xmls(f.filename, f.read()):
                try:
                    n = ler_nfe(conteudo)
                    notas.setdefault(n["chave"] or nome, n)  # mesma chave repetida no zip conta uma vez
                except ValueError as e:
                    ignorados.append({"arquivo": Path(nome).name, "motivo": str(e)})
        except zipfile.BadZipFile:
            ignorados.append({"arquivo": f.filename, "motivo": "zip inválido"})

    por_chave, _, erros = localizar_no_sismoura([n["chave"] for n in notas.values()], {})
    resumo = {"importadas": 0, "ja_lancadas": [], "duplicadas": 0, "ignorados": ignorados,
              "sem_loja": 0, "erros": erros}
    with local() as con:
        ja_existem = {r[0] for r in con.execute(
            "SELECT chave FROM notas_manuais WHERE chave IS NOT NULL AND resolvido_em IS NULL AND lancada_em IS NULL")}
        for n in notas.values():
            if n["chave"] in por_chave and por_chave[n["chave"]][0] <= dia:
                d, l = por_chave[n["chave"]]
                resumo["ja_lancadas"].append(f"NF {n['nf']} ({l}, {d[8:10]}/{d[5:7]})")
                continue
            if n["chave"] and n["chave"] in ja_existem:
                resumo["duplicadas"] += 1
                continue
            loja = loja_forcada or CNPJ_LOJA.get(n["cnpj_dest"])
            if not loja_ok(loja):
                ignorados.append({"arquivo": f"NF {n['nf']}", "motivo": f"destinada a {loja}"})
                continue
            if not loja and len(LOJAS) == 1:
                loja = next(iter(LOJAS.values()))
            if not loja:
                resumo["sem_loja"] += 1
            con.execute(
                """INSERT INTO notas_manuais (data, tipo, loja, nf, serie, fornecedor, valor, motivo,
                   chave, emissao, cnpj_emit, origem) VALUES (?,?,?,?,?,?,?,?,?,?,?, 'xml')""",
                (dia, tipo, loja, n["nf"], n["serie"], n["fornecedor"], n["valor"], motivo,
                 n["chave"] or None, n["emissao"], n["cnpj_emit"]))
            resumo["importadas"] += 1
    return jsonify(resumo)


def _data(s):
    """'2026-10-16' ou '16/10/2026' → date (ou o texto original, se não for data)."""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except (TypeError, ValueError):
            pass
    return s or None


def gerar_excel(p):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    AZUL, CINZA = "2F5BEA", "8A93A1"
    MOEDA, DATA = 'R$ #,##0.00;[Red]-R$ #,##0.00', "DD/MM/YYYY"
    titulo_f = Font(name="Calibri", size=18, bold=True, color="14171C")
    sub_f = Font(name="Calibri", size=11, color="4A5260")
    secao_f = Font(name="Calibri", size=12, bold=True, color=AZUL)
    cab_f = Font(bold=True, color="FFFFFF")
    cab_fill = PatternFill("solid", fgColor=AZUL)
    linha = Border(bottom=Side(style="thin", color="E3E6EB"))
    fechada_f = Font(italic=True, color=CINZA)
    dia = _data(p["data"])
    dia_txt = dia.strftime("%d/%m/%Y")
    nomes_tabela = iter(f"T{i}" for i in range(1, 100))

    wb = Workbook()

    def cabecalho(ws, titulo):
        ws["A1"] = titulo
        ws["A1"].font = titulo_f
        ws["A2"] = f"{dia_txt} · {p['loja']} · gerado em {datetime.now():%d/%m/%Y %H:%M}"
        ws["A2"].font = sub_f
        ws.sheet_view.showGridLines = False

    def tabela(ws, lin, colunas, linhas, formatos=None, total=None, apagadas=()):
        """Escreve uma tabela formatada a partir da linha `lin`. Retorna a próxima linha livre."""
        formatos = formatos or {}
        for c, nome in enumerate(colunas, 1):
            cel = ws.cell(lin, c, nome)
            cel.font, cel.fill = cab_f, cab_fill
            cel.alignment = Alignment(vertical="center")
        if not linhas:
            ws.cell(lin + 1, 1, "Nenhuma nota.").font = fechada_f
            return lin + 3
        for i, valores in enumerate(linhas, lin + 1):
            for c, v in enumerate(valores, 1):
                cel = ws.cell(i, c, v)
                cel.border = linha
                if c in formatos:
                    cel.number_format = formatos[c]
                if (i - lin - 1) in apagadas:
                    cel.font = fechada_f
        fim = lin + len(linhas)
        ref = f"A{lin}:{get_column_letter(len(colunas))}{fim}"
        t = Table(displayName=next(nomes_tabela), ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(t)
        if total:
            for c, v in total.items():
                cel = ws.cell(fim + 1, c, v)
                cel.font = Font(bold=True)
                if c in formatos:
                    cel.number_format = formatos[c]
            fim += 1
        return fim + 2

    def larguras(ws, valores):
        for c, w in enumerate(valores, 1):
            ws.column_dimensions[get_column_letter(c)].width = w

    # ---------- Resumo ----------
    ws = wb.active
    ws.title = "Resumo"
    cabecalho(ws, "Relatório de notas do dia")
    k = p["kpis"]
    ws.cell(4, 1, "Indicadores").font = secao_f
    lin = tabela(ws, 5, ["Indicador", "Qtde", "Valor", "Observação"], [
        ["Notas que chegaram hoje", k["chegaram"]["qt"], k["chegaram"]["valor"], "Lançadas + a lançar do dia"],
        ["Lançadas", k["lancadas"]["qt"], k["lancadas"]["valor"], f"por {len(p['usuarios'])} usuário(s)"],
        ["   · que chegaram hoje", k["lanc_hoje"]["qt"], k["lanc_hoje"]["valor"], "Entram em \"Notas que chegaram hoje\""],
        ["   · de dias anteriores", k["lanc_ant"]["qt"], k["lanc_ant"]["valor"], "Estavam a lançar/pendentes antes de hoje"],
        ["A lançar", k["alancar"]["qt"], k["alancar"]["valor"], "Chegaram hoje e ficaram para depois"],
        ["De dias anteriores", k["anteriores"]["qt"], k["anteriores"]["valor"],
         f"{k['anteriores']['mes_anterior']} do mês anterior"],
        ["Pendentes", k["pendentes"]["qt"], k["pendentes"]["valor"],
         f"{k['pendentes']['novas']} nova(s) hoje · {k['pendentes']['qt'] - k['pendentes']['novas']} anteriores"],
        ["Contas a pagar deste mês", k["cp_mes"]["qt"], k["cp_mes"]["valor"],
         f"Títulos das lançadas com vencimento até o fim de {k['cp_mes']['mes']}"],
        ["Contas a pagar dos próximos meses", k["cp_prox"]["qt"], k["cp_prox"]["valor"], k["cp_prox"]["meses"]],
        ["Bonificado (CFOP 1.910/2.910)", k["bonif"]["qt"], k["bonif"]["valor"],
         f"{k['bonif']['so_bonif']} só de bonificação · {k['bonif']['qt'] - k['bonif']['so_bonif']} mista(s)"],
    ], {3: MOEDA})
    ws.cell(5 + 1, 1).font = Font(bold=True)
    for r in (8, 9):  # detalhamento das lançadas
        ws.cell(r, 1).font = Font(color="4A5260")

    ws.cell(lin, 1, "Lançamentos por usuário").font = secao_f
    us = p["usuarios"]
    lin = tabela(ws, lin + 1, ["Usuário", "Notas", "Valor", "Lojas"],
                 [[u["nome"], u["qt"], u["valor"], u["lojas"]] for u in us], {3: MOEDA},
                 total={1: "Total", 2: sum(u["qt"] for u in us), 3: sum(u["valor"] for u in us)})

    ws.cell(lin, 1, "Resumo por loja").font = secao_f
    lj = p["lojas"]
    tabela(ws, lin + 1, ["Loja", "Lançadas", "A lançar", "Anteriores", "Pendentes", "Valor lançado"],
           [[r["loja"], r["l"], r["a"], r["n"], r["p"], r["v"]] for r in lj], {6: MOEDA},
           total={1: "Total", **{c: sum(r[key] for r in lj) for c, key in
                                 ((2, "l"), (3, "a"), (4, "n"), (5, "p"), (6, "v"))}})
    larguras(ws, [30, 11, 18, 40, 12, 18])

    # ---------- Lançadas ----------
    ws = wb.create_sheet("Lançadas")
    cabecalho(ws, "Notas lançadas no SISMOURA")
    la = p["lancadas"]
    tabela(ws, 4, ["Hora", "Loja", "NF", "Série", "Fornecedor", "Emissão", "Chegada", "Usuário", "Valor", "Bonificado"],
           [[n["hora"], n["loja"], n["nf"], n["serie"], n["fornecedor"], _data(n["emissao"]), _data(n["chegada"]),
             n["usuario"], n["valor"], n.get("bonificado") or 0] for n in la], {6: DATA, 7: DATA, 9: MOEDA, 10: MOEDA},
           total={8: f"Total ({len(la)} notas)", 9: sum(n["valor"] for n in la),
                  10: sum(n.get("bonificado") or 0 for n in la)})
    larguras(ws, [8, 20, 13, 7, 45, 12, 12, 30, 16, 16])
    ws.freeze_panes = "A5"

    # ---------- A lançar / Pendentes ----------
    for chave, titulo, com_motivo in (("alancar", "A lançar", False), ("pendentes", "Pendentes", True)):
        ws = wb.create_sheet(titulo)
        cabecalho(ws, f"Notas {titulo.lower()}")
        rows = p[chave]
        cols = ["Grupo", "Chegada", "Loja", "NF", "Série", "Fornecedor", "Emissão", "Valor"]
        cols += (["Motivo"] if com_motivo else []) + ["Situação", "Origem", "Chave de acesso"]
        dados = []
        for m in rows:
            d = [m["grupo"], _data(m["chegada"]), m["loja"], m["nf"], m["serie"], m["fornecedor"],
                 _data(m["emissao"]), m["valor"]]
            dados.append(d + ([m["motivo"]] if com_motivo else []) + [m["situacao"], m["origem"], m["chave"]])
        abertas = [m for m in rows if m["aberta"]]
        tabela(ws, 4, cols, dados, {2: DATA, 7: DATA, 8: MOEDA},
               total={6: f"Em aberto ({len(abertas)} notas)", 8: sum(m["valor"] for m in abertas)},
               apagadas={i for i, m in enumerate(rows) if not m["aberta"]})
        larguras(ws, [14, 12, 20, 13, 7, 45, 12, 16] + ([35] if com_motivo else []) + [24, 9, 48])
        ws.freeze_panes = "A5"

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@app.post("/api/excel")
def api_excel():
    p = request.get_json()
    return send_file(gerar_excel(p), as_attachment=True,
                     download_name=f"Relatorio de notas {p['data']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------- backup (exportar / importar as notas a lançar e pendentes) ----------
# As notas lançadas ficam no SISMOURA; aqui só vai o que existe apenas neste sistema.

SISTEMA_BACKUP = "relatorio-notas-por-dia"
CAMPOS_BACKUP = ("data", "tipo", "loja", "nf", "serie", "fornecedor", "valor", "motivo", "chave", "emissao",
                 "cnpj_emit", "origem", "resolvido_em", "lancada_em", "lancada_loja", "criado_em")


def notas_permitidas(con):
    return [r for r in con.execute("SELECT * FROM notas_manuais ORDER BY id") if loja_ok(r["loja"])]


def assinatura(n):
    """Identifica a mesma nota entre bancos, para não duplicar ao importar o backup de novo."""
    return (n["data"], n["tipo"], n["loja"] or "", so_digitos(n["nf"]), n["chave"] or "", n["criado_em"] or "")


def copiar_banco():
    """Cópia do banco local em dados/backups/ antes de importar. Retorna o nome do arquivo."""
    if not DB_LOCAL.exists():
        return None
    destino = DB_LOCAL.parent / "backups" / f"relatorio-{datetime.now():%Y%m%d-%H%M%S}.db"
    destino.parent.mkdir(exist_ok=True)
    orig, dest = sqlite3.connect(DB_LOCAL), sqlite3.connect(destino)
    try:
        orig.backup(dest)
    finally:
        orig.close()
        dest.close()
    return destino.name


@app.get("/api/backup/resumo")
def api_backup_resumo():
    with local() as con:
        rows = notas_permitidas(con)
    return jsonify({
        "total": len(rows),
        "alancar": sum(r["tipo"] == "alancar" for r in rows),
        "pendentes": sum(r["tipo"] == "pendente" for r in rows),
        "abertas": sum(not r["lancada_em"] and not r["resolvido_em"] for r in rows),
    })


@app.get("/api/backup")
def api_backup_exportar():
    with local() as con:
        notas = [{c: r[c] for c in CAMPOS_BACKUP} for r in notas_permitidas(con)]
    conteudo = {"sistema": SISTEMA_BACKUP, "versao": 1, "exportado_em": datetime.now().isoformat(timespec="seconds"),
                "lojas": list(LOJAS.values()), "notas": notas}
    buf = io.BytesIO(json.dumps(conteudo, ensure_ascii=False, indent=1).encode("utf-8"))
    return send_file(buf, as_attachment=True, mimetype="application/json",
                     download_name=f"backup-notas-{datetime.now():%Y-%m-%d-%H%M}.json")


@app.post("/api/backup")
def api_backup_importar():
    """Soma as notas do backup às atuais (nunca apaga). Repetidas e de outras lojas são ignoradas."""
    try:
        conteudo = json.load(request.files["arquivo"])
        notas = conteudo["notas"]
        if conteudo.get("sistema") != SISTEMA_BACKUP or not isinstance(notas, list):
            raise ValueError
    except (KeyError, ValueError, TypeError, UnicodeDecodeError):
        return jsonify({"erro": "O arquivo não é um backup deste relatório."}), 400

    resumo = {"importadas": 0, "repetidas": 0, "outra_loja": 0, "invalidas": 0, "copia": copiar_banco()}
    with local() as con:
        existentes = {assinatura(r) for r in con.execute("SELECT * FROM notas_manuais")}
        for n in notas:
            if not isinstance(n, dict) or not n.get("data") or n.get("tipo") not in ("alancar", "pendente"):
                resumo["invalidas"] += 1
                continue
            n = {c: n.get(c) for c in CAMPOS_BACKUP}
            if not loja_ok(n["loja"]):
                resumo["outra_loja"] += 1
                continue
            if assinatura(n) in existentes:
                resumo["repetidas"] += 1
                continue
            existentes.add(assinatura(n))
            n["valor"] = float(n["valor"] or 0)
            cols = [c for c in CAMPOS_BACKUP if n[c] is not None]
            con.execute(f"INSERT INTO notas_manuais ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                        [n[c] for c in cols])
            resumo["importadas"] += 1
    return jsonify(resumo)


@app.get("/")
def index():
    return send_from_directory(BASE / "web", "index.html")


PORTA_OCUPADA = 3  # código de saída usado por servidor/iniciar.ps1 para não ficar tentando de novo


def porta_em_uso(porta):
    with socket.socket() as s:
        s.settimeout(1)
        return s.connect_ex(("127.0.0.1", porta)) == 0


if __name__ == "__main__":
    # Uso:  python app.py                 → modo local: só este computador acessa, abre o navegador
    #       python app.py --producao      → servidor da empresa: Waitress, aceita acesso da rede
    #                                        (http://NOME-DO-SERVIDOR:3003), como o Controle de validade
    producao = "--producao" in sys.argv
    host = "0.0.0.0" if producao else "127.0.0.1"
    if porta_em_uso(PORTA):
        print(f"ERRO: a porta {PORTA} já está em uso. Veja qual programa a ocupa antes de iniciar "
              f"(não troque de porta sem atualizar orientacoes-projetos-ia.md).", file=sys.stderr)
        sys.exit(PORTA_OCUPADA)
    print(f"Lojas liberadas neste servidor: {', '.join(LOJAS.values())}")
    if producao:
        print(f"Escutando na rede: http://{socket.gethostname()}:{PORTA} (produção)", flush=True)
        from waitress import serve
        serve(app, host=host, port=PORTA, threads=8)
    else:
        print(f"Escutando em http://localhost:{PORTA} (local)", flush=True)
        if "--sem-navegador" not in sys.argv:
            threading.Timer(1.2, lambda: webbrowser.open(f"http://localhost:{PORTA}")).start()
        app.run(host=host, port=PORTA, debug=False)
